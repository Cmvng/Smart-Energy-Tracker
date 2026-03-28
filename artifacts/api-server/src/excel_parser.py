import sys, json, re, os
from datetime import datetime

def parse_date(s):
    if not s: return None
    s = str(s).strip().split(' ')[0]
    m = re.match(r'^(\d{1,2})[/\-](\d{1,2})[/\-](\d{2,4})$', s)
    if m:
        day, month, year_raw = m.groups()
        yr = int(year_raw)
        year = str(2000+yr) if yr<100 and yr<50 else (
               str(1900+yr) if yr<100 else str(yr))
        try:
            dt = datetime.strptime(
                f'{year}-{month.zfill(2)}-{day.zfill(2)}',
                '%Y-%m-%d')
            if 2000 <= dt.year <= 2030:
                return dt.strftime('%Y-%m-%d')
        except: pass
    m = re.match(r'^(\d{4})[/\-](\d{1,2})[/\-](\d{1,2})$', s)
    if m:
        y, mo, d = m.groups()
        try:
            dt = datetime.strptime(
                f'{y}-{mo.zfill(2)}-{d.zfill(2)}',
                '%Y-%m-%d')
            if 2000 <= dt.year <= 2030:
                return dt.strftime('%Y-%m-%d')
        except: pass
    months_map = {
        'jan':'01','feb':'02','mar':'03','apr':'04',
        'may':'05','jun':'06','jul':'07','aug':'08',
        'sep':'09','oct':'10','nov':'11','dec':'12'
    }
    m = re.search(
        r'(\d{1,2})\s+(Jan|Feb|Mar|Apr|May|Jun|Jul|'
        r'Aug|Sep|Oct|Nov|Dec)\w*\s+(\d{4})',
        s, re.IGNORECASE)
    if m:
        d, mn, y = m.groups()
        mo = months_map[mn.lower()[:3]]
        try:
            dt = datetime.strptime(
                f'{y}-{mo}-{d.zfill(2)}', '%Y-%m-%d')
            if 2000 <= dt.year <= 2030:
                return dt.strftime('%Y-%m-%d')
        except: pass
    return None

def clean_amount(s):
    if not s: return None
    s = str(s).strip()
    is_neg = s.startswith('(') or s.startswith('-')
    s = re.sub(r'[₦£\$€,\s()A-Za-z]', '', s)
    if not s: return None
    try:
        v = float(s)
        return abs(v) if v != 0 else None
    except: return None

def is_negative_amount(s):
    if not s: return False
    s = str(s).strip()
    return s.startswith('(') or s.startswith('-') or \
           s.upper().endswith('DR')

INCOME_KEYWORDS = [
    'inward','credit','reversal','refund','received',
    'money in','salary','deposit','lodg','inflow',
    'transfer in','nip in','cr ','nibss credit',
    'mpesa credit','momo credit','direct deposit',
    'ach credit','bacs credit','faster payment in',
]
EXPENSE_KEYWORDS = [
    'outward','debit','web payment','atm','pos',
    'charge','fee','stamp','bill','subscription',
    'payment','purchase','withdrawal','transfer out',
    'nip out',' dr ','mpesa debit','momo debit',
    'direct debit','ach debit',
]
SKIP_KEYWORDS = [
    'opening balance','closing balance','brought forward',
    'carried forward','available balance','ledger balance',
    'total debit','total credit','account balance',
    'b/f','c/f',
]
FINANCIAL_KEYWORDS = [
    'date','money in','money out','debit','credit',
    'amount','narration','description','withdrawal',
    'deposit','particulars','transaction','balance',
    'trans date','value date','dr','cr','ledger',
    'reference','details','remarks','beneficiary',
]

def detect_type(cat, desc, has_in, has_out, is_neg):
    if has_in and not has_out: return 'income'
    if has_out and not has_in: return 'expense'
    text = (str(cat)+' '+str(desc)).lower()
    for k in SKIP_KEYWORDS:
        if k in text: return None
    for k in INCOME_KEYWORDS:
        if k in text: return 'income'
    for k in EXPENSE_KEYWORDS:
        if k in text: return 'expense'
    if is_neg: return 'expense'
    return 'income'

def detect_currency(rows):
    text = ' '.join(str(c) for row in rows[:20]
                   for c in row if c)
    scores = {
        'NGN': len(re.findall(r'₦|NGN|naira', text, re.I)),
        'GBP': len(re.findall(r'£|GBP|pound', text, re.I)),
        'USD': len(re.findall(r'\$|USD|dollar', text, re.I)),
        'EUR': len(re.findall(r'€|EUR|euro', text, re.I)),
        'KES': len(re.findall(r'KSh|KES|shilling', text, re.I)),
        'GHS': len(re.findall(r'₵|GHS|cedi', text, re.I)),
        'ZAR': len(re.findall(r'ZAR|rand', text, re.I)),
    }
    best = max(scores, key=scores.get)
    return best if scores[best] > 0 else 'USD'

def score_row(row):
    cells = [str(c or '').lower().strip() for c in row if c]
    return sum(1 for c in cells
               if any(kw in c for kw in FINANCIAL_KEYWORDS))

def find_best_sheet_and_header(workbook):
    best_sheet_rows = None
    best_score = 0
    
    for sheet_name in workbook.sheetnames:
        try:
            ws = workbook[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            for i, row in enumerate(rows[:60]):
                s = score_row(row)
                if s > best_score:
                    best_score = s
                    best_sheet_rows = (rows, i)
        except: continue
    
    return best_sheet_rows, best_score

def map_columns(headers):
    colDate=colIn=colOut=colDesc=colCat=colBal=colAmt=-1
    for i, h in enumerate(headers):
        h = h.lower().strip()
        if ('balance' in h or h in ['bal','running bal']) \
           and colBal == -1:
            colBal = i; continue
        if (('money in' in h or h in ['credit','cr','deposit'] or
             'paid in' in h or 'inflow' in h or
             'lodg' in h or 'cr amount' in h or
             'credit amount' in h or 'deposits' in h) and
            'debit' not in h and 'descri' not in h and
            colIn == -1): colIn = i
        if (('money out' in h or h in ['debit','dr','withdrawal'] or
             'paid out' in h or 'outflow' in h or
             'dr amount' in h or 'debit amount' in h or
             'withdrawals' in h) and
            'credit' not in h and colOut == -1): colOut = i
        if (('date' in h or h == 'date/time' or
             'trans date' in h or 'value date' in h or
             'posting date' in h) and
            colDate == -1): colDate = i
        if (('descri' in h or 'narrat' in h or
             'partic' in h or 'detail' in h or
             'to / from' in h or 'to/from' in h or
             'remark' in h or 'benefi' in h or
             'merchant' in h or 'memo' in h or
             h == 'particulars') and
            colDesc == -1): colDesc = i
        if 'categ' in h and colCat == -1: colCat = i
        if h in ['amount','transaction amount','value'] \
           and colAmt == -1: colAmt = i
    return colDate,colIn,colOut,colDesc,colCat,colBal,colAmt

def parse_excel(filepath):
    import openpyxl
    
    try:
        wb = openpyxl.load_workbook(
            filepath, read_only=True, data_only=True)
    except Exception as e:
        print(json.dumps({
            'success': False,
            'error': f'Cannot open file: {str(e)}',
            'transactions': [], 'total_found': 0
        }))
        return
    
    result, score = find_best_sheet_and_header(wb)
    wb.close()
    
    if not result or score < 2:
        try:
            wb2 = openpyxl.load_workbook(
                filepath, read_only=True, data_only=True)
            ws = wb2.worksheets[0]
            rows = list(ws.iter_rows(values_only=True))
            wb2.close()
            header_idx = 0
            for i, row in enumerate(rows):
                if any(c for c in row):
                    header_idx = i
                    break
            result = (rows, header_idx)
            score = 1
        except:
            print(json.dumps({
                'success': False,
                'error': 'Could not read file',
                'transactions': [], 'total_found': 0
            }))
            return
    
    rows, header_idx = result
    if not rows:
        print(json.dumps({
            'success': False,
            'error': 'Empty file',
            'transactions': [], 'total_found': 0
        }))
        return
    
    headers = [str(c or '').strip() 
               for c in rows[header_idx]]
    colDate,colIn,colOut,colDesc,colCat,colBal,colAmt = \
        map_columns(headers)
    
    if colDate == -1: colDate = 0
    
    if colDesc == -1:
        excl = {colDate,colIn,colOut,colBal,colAmt,colCat}
        max_len = 0
        for ci in range(len(headers)):
            if ci in excl: continue
            avg = sum(
                len(str(rows[ri][ci] or ''))
                for ri in range(
                    header_idx+1,
                    min(header_idx+6, len(rows))
                )
                if ci < len(rows[ri])
            ) / 5
            if avg > max_len:
                max_len = avg
                colDesc = ci
    
    currency = detect_currency(rows)
    transactions = []
    
    for row in rows[header_idx+1:]:
        if not row or not any(c for c in row): continue
        
        g = lambda idx: (
            row[idx] if idx >= 0 and idx < len(row) 
            else None
        )
        
        date_raw = str(g(colDate) or '').strip()
        if not date_raw or date_raw.lower() in [
            'none','nan','','-']: continue
        date = parse_date(date_raw.split(' ')[0])
        if not date: continue
        
        raw_in = g(colIn) if colIn >= 0 else None
        raw_out = g(colOut) if colOut >= 0 else None
        
        in_amt = clean_amount(raw_in)
        out_amt = clean_amount(raw_out)
        
        if colAmt >= 0 and not in_amt and not out_amt:
            raw_amt = g(colAmt)
            av = clean_amount(raw_amt)
            if av:
                if is_negative_amount(raw_amt):
                    out_amt = av
                else:
                    in_amt = av
        
        if not in_amt and not out_amt: continue
        
        desc = str(g(colDesc) or g(colCat) or 
                   'Transaction').strip()[:80]
        cat = str(g(colCat) or '').strip()
        
        if any(k in desc.lower() for k in SKIP_KEYWORDS):
            continue
        if any(k in cat.lower() for k in SKIP_KEYWORDS):
            continue
        
        is_neg = is_negative_amount(
            g(colOut) if colOut >= 0 else g(colAmt)
        )
        
        t = detect_type(
            cat, desc,
            bool(in_amt and in_amt > 0),
            bool(out_amt and out_amt > 0),
            is_neg
        )
        if not t: continue
        
        amount = in_amt if t == 'income' else out_amt
        if not amount or amount <= 0: continue
        
        transactions.append({
            'date': date,
            'description': desc,
            'amount': amount,
            'type': t,
            'currency': currency,
            'confidence': 'high'
        })
    
    print(json.dumps({
        'success': True,
        'method': 'excel-python',
        'currency': currency,
        'total_found': len(transactions),
        'transactions': transactions
    }))

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(json.dumps({'error': 'No file provided'}))
        sys.exit(1)
    parse_excel(sys.argv[1])
